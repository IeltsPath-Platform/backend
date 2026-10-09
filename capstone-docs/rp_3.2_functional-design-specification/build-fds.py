"""Build the IELTSPath FDS from the FDS template and the content in fds_content.py.

The template's repeatable blocks (one screen section, one background-job section) are cloned with their
original formatting; only text is replaced. Placeholder runs (italic grey) become regular black text.

Usage: python build-fds.py [output.docx]   (writes Report 3.2_FDS_IELTSPath_v1.docx next to this script by default)
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Cm

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import fds_content as C  # noqa: E402

TEMPLATE = HERE / "Report 3.2_FDS_Template.docx"
OUTPUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Report 3.2_FDS_IELTSPath_v1.docx"
PLACEHOLDER_GREY = "595959"
CODE_FONT = "Courier New"

doc = docx.Document(str(TEMPLATE))
body = doc.element.body
E = list(body.iterchildren())
if len(E) != 166 or E[165].tag != qn("w:sectPr"):
    raise SystemExit(f"Unexpected template layout ({len(E)} body elements); update the indices in build-fds.py")

# Prototype elements, addressed by their position in the template body.
P = {
    "h1": E[35], "h2": E[38], "h3": E[39], "h4": E[43], "note": E[40], "code": E[72], "spacer": E[42],
    "label": E[121], "label_value": E[145], "text": E[98], "end_note": E[164],
}
T = {
    "doc_info": E[4], "versions": E[8], "kv": E[41], "nav": E[44], "pre": E[51], "entry": E[55], "exit": E[59],
    "comp": E[63], "api": E[67], "msg": E[81], "rule": E[85], "job_kv": E[93], "input": E[122],
    "output": E[125], "job_rule": E[139], "error": E[143], "sla": E[149], "log": E[156], "alert": E[161],
}


# ---------------------------------------------------------------- run / text helpers
def _normalise_rpr(rpr):
    """Turns a placeholder run (italic grey) into regular black text; other styling is kept."""
    if rpr is None:
        return rpr
    color = rpr.find(qn("w:color"))
    if color is not None and color.get(qn("w:val")) == PLACEHOLDER_GREY:
        color.set(qn("w:val"), "000000")
        for tag in ("w:i", "w:iCs"):
            el = rpr.find(qn(tag))
            if el is not None:
                rpr.remove(el)
    return rpr


def _code_rpr(rpr):
    rpr = copy.deepcopy(rpr) if rpr is not None else docx.oxml.OxmlElement("w:rPr")
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = docx.oxml.OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), CODE_FONT)
    return rpr


def _run(rpr, text):
    r = docx.oxml.OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = docx.oxml.OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def _fill_paragraph(p, text, rpr, normalise=True, code_spans=True):
    """Replaces all runs of paragraph p by text; `inline code` spans use the code font."""
    for r in list(p):
        if r.tag != qn("w:pPr"):
            p.remove(r)
    base = _normalise_rpr(copy.deepcopy(rpr)) if normalise else copy.deepcopy(rpr)
    parts = re.split(r"`([^`]+)`", text) if code_spans else [text]
    for i, part in enumerate(parts):
        if part:
            p.append(_run(_code_rpr(base) if i % 2 else base, part))


def _first_rpr(el):
    r = el.find(".//" + qn("w:r"))
    return r.find(qn("w:rPr")) if r is not None else None


# ---------------------------------------------------------------- paragraph builders
def para(kind, text, normalise=True):
    p = copy.deepcopy(P[kind])
    _fill_paragraph(p, text, _first_rpr(P[kind]), normalise=normalise, code_spans=kind != "code")
    return p


def code_block(lines):
    return [para("code", line if line else " ") for line in lines]


def label_value(label, value):
    p = copy.deepcopy(P["label_value"])
    runs = p.findall(qn("w:r"))
    label_rpr, value_rpr = runs[0].find(qn("w:rPr")), runs[1].find(qn("w:rPr"))
    for r in runs:
        p.remove(r)
    p.append(_run(label_rpr, label + " "))
    tmp = docx.oxml.OxmlElement("w:p")
    _fill_paragraph(tmp, value, value_rpr)
    for r in tmp.findall(qn("w:r")):
        p.append(r)
    return p


def picture(path, width_cm=16.5):
    holder = doc.add_paragraph()
    holder.alignment = 1
    holder.add_run().add_picture(str(path), width=Cm(width_cm))
    el = holder._p
    body.remove(el)
    return el


def caption(text):
    p = copy.deepcopy(P["end_note"])  # centred, italic grey, small
    _fill_paragraph(p, text, _first_rpr(P["end_note"]), normalise=False)
    return p


# ---------------------------------------------------------------- table builders
def _fill_cell(tc, text):
    paras = tc.findall(qn("w:p"))
    proto = paras[0]
    rpr = _first_rpr(proto)
    for extra in paras[1:]:
        tc.remove(extra)
    lines = str(text).split("\n")
    _fill_paragraph(proto, lines[0], rpr)
    anchor = proto
    for line in lines[1:]:
        p = copy.deepcopy(proto)
        _fill_paragraph(p, line, rpr)
        anchor.addnext(p)
        anchor = p


def _cell_text(tc):
    return "".join(t.text or "" for t in tc.iter(qn("w:t"))).strip()


# Column widths (dxa, total 9026 like the template) where long ids or code paths need more room than the
# template grid gives; every table also gets a fixed layout so Word does not autofit columns to code text.
WIDTHS = {
    "pre": [800, 2200, 1900, 4126],
    "entry": [800, 2700, 2700, 2826],
    "exit": [800, 2100, 2000, 2000, 2126],
    "comp": [2300, 1700, 1500, 2100, 1426],
    "api": [380, 1400, 820, 1900, 1500, 1500, 1526],
}
CATALOG_WIDTHS = [820, 3000, 900, 1400, 1000, 800, 1106]


def _fix_layout(tbl, widths=None):
    tbl_pr = tbl.find(qn("w:tblPr"))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = docx.oxml.OxmlElement("w:tblLayout")
        # Schema order: … tblBorders, shd, tblLayout, tblCellMar, tblLook — Word repairs the file otherwise.
        after = next((tbl_pr.find(qn(t)) for t in ("w:tblCellMar", "w:tblLook") if tbl_pr.find(qn(t)) is not None),
                     None)
        if after is not None:
            after.addprevious(layout)
        else:
            tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    if not widths:
        return
    for col, w in zip(tbl.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w))
    for tr in tbl.findall(qn("w:tr")):
        for tc, w in zip(tr.findall(qn("w:tc")), widths):
            tcw = tc.find(qn("w:tcPr")).find(qn("w:tcW"))
            tcw.set(qn("w:w"), str(w))
            tcw.set(qn("w:type"), "dxa")


def table(kind, rows, header=None, match_col=None, widths=None):
    tbl = _table(kind, rows, header, match_col)
    _fix_layout(tbl, widths or WIDTHS.get(kind))
    return tbl


def _table(kind, rows, header=None, match_col=None):
    """Clones a template table: the header row is kept (or relabelled by `header`) and every data row is
    cloned from a template data row. With match_col, the template row whose cell in that column has the
    same text as the new row is used, so row colouring follows the row type (e.g. Transient / Critical)."""
    tbl = copy.deepcopy(T[kind])
    trs = tbl.findall(qn("w:tr"))
    head, protos = trs[0], trs[1:]
    if header:
        for tc, value in zip(head.findall(qn("w:tc")), header):
            _fill_cell(tc, value)
    keyed = {}
    if match_col is not None:
        for idx, tr in enumerate(protos):
            keyed.setdefault(_cell_text(tr.findall(qn("w:tc"))[match_col]), idx)
    for tr in protos:
        tbl.remove(tr)
    for row in rows:
        idx = keyed.get(row[match_col], 0) if match_col is not None else 0
        tr = copy.deepcopy(protos[idx])
        for tc, value in zip(tr.findall(qn("w:tc")), row):
            _fill_cell(tc, value)
        tbl.append(tr)
    return tbl


def kv_table(kind, rows):
    """Two-column key/value tables have no header row; rows alternate the two template shadings."""
    tbl = copy.deepcopy(T[kind])
    trs = tbl.findall(qn("w:tr"))
    protos = trs[:2]
    for tr in trs:
        tbl.remove(tr)
    for i, (key, value) in enumerate(rows):
        tr = copy.deepcopy(protos[i % 2])
        tcs = tr.findall(qn("w:tc"))
        _fill_cell(tcs[0], key)
        _fill_cell(tcs[1], value)
        tbl.append(tr)
    _fix_layout(tbl)
    return tbl


# ---------------------------------------------------------------- sections
out = []
add = out.extend


def spacer():
    return copy.deepcopy(P["spacer"])


def build_front():
    title_block = [copy.deepcopy(e) for e in E[0:4]]
    add(title_block)
    add([kv_table("doc_info", C.DOC_INFO), spacer(), para("h1", "Version History"),
         table("versions", C.VERSION_HISTORY), spacer()])


def build_part1():
    add([para("h1", "PART 1 – OVERVIEW")])
    add([para("h2", "1.1  Screen Navigation Flow")])
    for line in C.NAV_FLOW_TEXT:
        add([para("text", line)])
    add([picture(C.NAV_FLOW_PNG), caption(C.NAV_FLOW_CAPTION)])
    add([para("h4", "Screen Index"),
         table("exit", C.screen_index_rows(), header=("Section", "Screen", "Route", "Role(s)", "Status")),
         spacer()])
    add([para("h2", "1.2  Job Schedule & Dependencies")])
    for line in C.JOB_FLOW_TEXT:
        add([para("text", line)])
    add([picture(C.JOB_FLOW_PNG), caption(C.JOB_FLOW_CAPTION)])
    add(code_block(C.JOB_SCHEDULE_ASCII))
    add([spacer()])


def build_conventions():
    add([para("h1", "PART 2 – SCREEN SPECIFICATIONS")])
    for line in C.PART2_INTRO:
        add([para("text", line)])
    add([para("h2", "2.0  Common Screen Conventions")])
    for line in C.CONVENTIONS_TEXT:
        add([para("text", line)])
    add([para("h4", "Global Messages"), table("msg", C.GLOBAL_MESSAGES), spacer()])
    add([para("h4", "Global Guard Redirects"), table("pre", C.GLOBAL_GUARDS), spacer()])


MISSING_SHOTS = []


def _shots(s):
    """Screenshots of a screen whose files exist; a missing file is reported and left out of the document."""
    found = []
    for file, text in s.get("shots", []):
        if (C.SCREENSHOT_DIR / file).exists():
            found.append((file, text))
        else:
            MISSING_SHOTS.append(file)
    return found


def _mockup_text(num, s, shots):
    if shots:
        figures = ", ".join(f"Figure {num}-{i}" for i in range(1, len(shots) + 1))
        return f"{figures} below — screenshot of the frontend (IELTS Space web app)"
    return s.get("mockup", "TBD — no frontend screen yet")


def build_screen(n, s):
    num = f"2.{n}"
    key = s["key"]
    add([para("h2", f"{num}  {s['name']}")])
    if s.get("draft"):
        add([para("note", s["draft"], normalise=False)])
    add([para("h3", f"{num}.1  General Information")])
    shots = _shots(s)
    info = [
        ("Screen Name", s["name"]),
        ("Purpose", s["purpose"]),
        ("Route / Type", s["route"]),
        ("Role(s)", s["roles"]),
        ("Related FT / UC", f"{s['ft']}\n{C.uc_label(s['uc'])}"),
        ("Status", s["status"]),
        ("Frontend", s.get("frontend", C.FE_NOT_BUILT)),
        ("Mockup / Wireframe", _mockup_text(num, s, shots)),
    ]
    if C.open_questions_for(s["name"]):
        info.append(("Open Questions", C.open_questions_for(s["name"]) + " (Appendix B)"))
    add([kv_table("kv", info), spacer()])
    for i, (file, text) in enumerate(shots, 1):
        add([picture(C.SCREENSHOT_DIR / file, width_cm=15.5),
             caption(f"Figure {num}-{i} — {text}. Source: screenshots/{file}")])
    if shots:
        add([spacer()])
    add([para("h4", "Navigation Context"),
         table("nav", s["nav_from"]), spacer(),
         table("nav", s["nav_to"], header=("Navigates to", "Condition / Trigger")), spacer()])

    add([para("h3", f"{num}.2  Entry & Exit Conditions"), para("h4", "Preconditions & Guard Redirects")])
    add([table("pre", [(f"PC-{i:02d}", *r) for i, r in enumerate(s["pre"], 1)]), spacer()])
    add([para("h4", "Entry Triggers"),
         table("entry", [(f"ET-{i:02d}", *r) for i, r in enumerate(s["entry"], 1)]), spacer()])
    add([para("h4", "Exit Points"),
         table("exit", [(f"EX-{i:02d}", *r) for i, r in enumerate(s["exits"], 1)]), spacer()])

    add([para("h3", f"{num}.3  UI Components")])
    comps = []
    counters = {}
    for ctype, name, kind, validation, notes in s["comps"]:
        counters[ctype] = counters.get(ctype, 0) + 1
        comps.append((f"{key}-{ctype}-{counters[ctype]:02d}", name, kind, validation, notes))
    add([table("comp", comps), spacer()])

    add([para("h3", f"{num}.4  API Calls & Data Flow")])
    rows = []
    for i, (trigger, api_key, request, success, error) in enumerate(s["apis"], 1):
        if api_key is None:  # screen built on mock data or browser storage
            rows.append((str(i), trigger, "—", "No API (mock data)", request, success, error))
            continue
        api_id, method_path = C.api(api_key)
        C.register_use(api_key, num)
        rows.append((str(i), trigger, api_id, f"`{method_path}`", request, success, error))
    add([table("api", rows), spacer()])

    add([para("h3", f"{num}.5  Interaction Specifications")])
    for title, lines in s["interactions"]:
        add([para("h4", title)])
        add(code_block(lines))
    add([spacer()])

    add([para("h3", f"{num}.6  Messages & Display Content")])
    add([table("msg", [(f"MSG-{i:02d}", *r) for i, r in enumerate(s["msgs"], 1)]), spacer()])

    add([para("h3", f"{num}.7  Business Rules")])
    add([table("rule", s["rules"]), spacer()])


def build_jobs():
    add([para("h1", "PART 3 – BACKGROUND JOB SPECIFICATIONS")])
    for line in C.PART3_INTRO:
        add([para("text", line)])
    for n, j in enumerate(C.JOBS, 1):
        num = f"3.{n}"
        add([para("h2", f"{num}  [{j['code']}] {j['name']}")])
        if j.get("draft"):
            add([para("note", j["draft"], normalise=False)])
        add([para("h3", f"{num}.1  General Information")])
        add([kv_table("job_kv", [
            ("Code / Name", f"{j['code']} — {j['name']}"),
            ("Class / Bean", j["bean"]),
            ("Purpose", j["purpose"]),
            ("Trigger / Schedule", j["trigger"]),
            ("Related FT / BR", j["refs"]),
            ("Status", j["status"]),
            ("SLA (max runtime)", j["sla_short"]),
        ] + ([("Open Questions", C.open_questions_for(j["code"]) + " (Appendix B)")]
             if C.open_questions_for(j["code"]) else [])), spacer()])
        add([para("h4", "Business Context"), para("text", j["context"]), spacer()])
        add([para("h3", f"{num}.2  Detailed Specification"), para("h4", "Processing Flow")])
        add(code_block(j["flow"]))
        add([spacer(), para("h4", "Input / Output Data"), para("label", "Input:"),
             table("input", j["inputs"]), spacer(), para("label", "Output:"),
             table("output", j["outputs"]), spacer()])
        add([para("h4", "Idempotency Guarantee")])
        add(code_block(j["idempotency"]))
        add([label_value("Verification:", j["verification"]), spacer()])
        add([para("h4", "Business Rules"), table("job_rule", j["rules"]), spacer()])
        add([para("h4", "Error Handling & Retry Policy"),
             table("error", j["errors"], match_col=0), spacer(),
             label_value("Retry policy:", j["retry"]), label_value("DLQ (if event-driven):", j["dlq"]), spacer()])
        add([para("h4", "SLA / Performance Expectation"), table("sla", j["sla"]), spacer(),
             label_value("Required index:", j["index"]), label_value("Batch strategy:", j["batch"]), spacer()])
        add([para("h4", "Logging & Alerting"), para("label", "Required log events:"),
             table("log", j["logs"]), spacer(), label_value("Log format:", j["log_format"]), spacer(),
             para("label", "Alert rules:"), table("alert", j["alerts"], match_col=2), spacer()])


def build_appendix():
    add([para("h1", "APPENDIX A – API CATALOG")])
    for line in C.APPENDIX_INTRO:
        add([para("text", line)])
    add([table("api", C.api_catalog_rows(),
               header=("API ID", "Method + Path", "Service", "Role(s)", "Screens", "FT", "Status"),
               widths=CATALOG_WIDTHS), spacer()])
    add([para("h1", "APPENDIX B – OPEN QUESTIONS")])
    add([table("rule", C.OPEN_QUESTIONS, header=("#", "Open question", "Affects")), spacer()])
    add([para("h1", "APPENDIX C – FRONTEND ALIGNMENT GAPS")])
    add([para("text", "Differences between the frontend (" + C.FRONTEND_REF + ") and the SRS / this FDS, found "
                      "while reading the frontend code. Each gap is also noted in the affected screen section.")])
    add([table("exit", C.FE_GAPS, header=("#", "Gap", "Screen", "SRS ref", "Suggested action"),
               widths=[900, 3600, 1600, 1300, 1626]), spacer()])
    build_traceability()


def _traceability_rows():
    """Use case → screens and business rule → screens/jobs, read from the screen and job definitions."""
    uc_screens = {uc: [] for uc in C.USE_CASES}
    br_users = {br: [] for br in C.BUSINESS_RULES}
    for n, s in enumerate(C.SCREENS, 1):
        label = f"2.{n} {s['name']}"
        for uc in C.uc_ids(s["uc"]):
            uc_screens[uc].append(label)
        for br in sorted({b for r in s["rules"] for b in re.findall(r"BR-\d{2}", r[0])}):
            br_users.setdefault(br, []).append(f"2.{n}")
    for j in C.JOBS:
        refs = " ".join([j["refs"]] + [r[0] for r in j["rules"]])
        for br in sorted(set(re.findall(r"BR-\d{2}", refs))):
            br_users.setdefault(br, []).append(j["code"])
    unknown = sorted(set(br_users) - set(C.BUSINESS_RULES))
    if unknown:
        raise SystemExit(f"Business rules {unknown} are not in the RTW Business Rules sheet (fds_content.BUSINESS_RULES)")
    uc_rows = [(uc + (" *" if uc in C.PROPOSED_UC_IDS else ""), name, ft, status,
                "\n".join(uc_screens[uc]) or "No screen — gap")
               for uc, (name, ft, status) in C.USE_CASES.items()]
    br_rows = [(br, title, status, owner, ", ".join(dict.fromkeys(br_users[br])) or "Not referenced — gap")
               for br, (title, status, owner) in C.BUSINESS_RULES.items()]
    return uc_rows, br_rows


def build_traceability():
    uc_rows, br_rows = _traceability_rows()
    add([para("h1", "APPENDIX D – TRACEABILITY TO THE RTW")])
    add([para("text", "Links every use case (RTW Sheet 2) and business rule (RTW Sheet 6) of " + C.RTW_REF + " to "
                      "the screens of Part 2 and the jobs of Part 3 that realise it. The tables are generated from the "
                      "screen and job sections, so they always match the body of this document. Status is the status "
                      "in SRS v0.9.24: Specified / Approved = in effect in the current release; Draft = required but "
                      "not built yet. Use cases marked * and BR-35 to BR-41 come from the SRS and are not in RTW v1.0 "
                      "yet; their UC ids are proposed (OQ-21). \"gap\" marks an item that no screen or job covers.")])
    add([para("h4", "D.1  Use Case → Screens"),
         table("exit", uc_rows, header=("UC ID", "Use case", "FT", "Status", "Screen(s)"),
               widths=[800, 2700, 1700, 1100, 2726]), spacer()])
    add([para("h4", "D.2  Business Rule → Screens and Jobs"),
         table("exit", br_rows, header=("BR ID", "Rule (short)", "Status", "Owner", "Used in"),
               widths=[800, 3000, 1100, 1900, 2226]), spacer()])
    gaps = [r[0] for r in uc_rows if r[4].endswith("gap")] + [r[0] for r in br_rows if r[4].endswith("gap")]
    print("Traceability gaps:", ", ".join(gaps) or "none")


def build_tail():
    add([copy.deepcopy(E[163]), copy.deepcopy(E[164])])


build_front()
build_part1()
build_conventions()
for n, screen in enumerate(C.SCREENS, 1):
    build_screen(n, screen)
build_jobs()
build_appendix()
build_tail()

sect = E[165]
for el in list(body.iterchildren()):
    body.remove(el)
for el in out:
    body.append(el)
body.append(sect)

# Pictures are created one by one and moved, so python-docx gives each the same drawing id; Word needs unique ids.
for n, doc_pr in enumerate(body.iter("{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"),
                           1):
    doc_pr.set("id", str(n))

# Text that starts or ends with a space (template runs included) keeps it only with xml:space="preserve".
for t in body.iter(qn("w:t")):
    if t.text and t.text != t.text.strip():
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

# Running header: project name and version.
for section in doc.sections:
    for p in section.header.paragraphs:
        for r in p.runs:
            r.text = r.text.replace("[Project Name]", "IELTSPath").replace("v[X.Y]", C.VERSION)
    full = "".join(r.text for p in section.header.paragraphs for r in p.runs)
    if "[" in full:
        for p in section.header.paragraphs:
            text = "".join(r.text for r in p.runs)
            text = text.replace("[Project Name]", "IELTSPath").replace("v[X.Y]", C.VERSION)
            text = re.sub(r"v\[X\.Y\]|\[X\.Y\]", C.VERSION, text)
            for r in p.runs[1:]:
                r.text = ""
            if p.runs:
                p.runs[0].text = text

doc.core_properties.title = "IELTSPath — Functional Design Specification"
doc.core_properties.author = C.AUTHOR
doc.save(str(OUTPUT))
print(f"Wrote {OUTPUT.name}: {len(C.SCREENS)} screens, {len(C.JOBS)} jobs")
if MISSING_SHOTS:
    print("Screenshots not found (left out):", ", ".join(MISSING_SHOTS))
