"""Replace only SRS section 4.1; preserve other OOXML package parts."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from io import BytesIO
import json, hashlib
from lxml import etree as E
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

DIR=Path(__file__).resolve().parent
SRS=DIR.parent/'Report 3.0_SRS_IELTSPath_v1.docx'
BACKUP=DIR/'qa'/'original-srs.docx'
BACKUP.parent.mkdir(exist_ok=True)
if not BACKUP.exists():BACKUP.write_bytes(SRS.read_bytes())
MODEL=json.loads((DIR/'erd-model.json').read_text())
VIEWS=json.loads((DIR/'diagram-manifest.json').read_text())
N=MODEL['nodes']
doc=Document(BACKUP)
body=doc._element.body
text=lambda el:''.join(el.itertext()) if False else ''.join(el.xpath('.//w:t/text()'))
children=list(body)
start=next(i for i,e in enumerate(children) if text(e).startswith('4.1  Core Entities'))
end=next(i for i,e in enumerate(children) if text(e).startswith('4.2  Data Constraints'))
placeholder=next(e for e in children[start:end] if e.tag==qn('w:tbl') and text(e).startswith('Entity'))
table_template=deepcopy(placeholder)
for el in children[start+1:end]:body.remove(el)
anchor=children[end]
new_elements=[]

def para(value='',style=None,bold=False,break_before=False):
    p=doc.add_paragraph(value,style)
    if bold:
        for r in p.runs:r.bold=True
    if break_before:p.paragraph_format.page_break_before=True
    anchor.addprevious(p._p); new_elements.append(p._p)
    return p

def relationship_text(key):
    t=N[key]; out=[]
    # Describe ownership from the perspective of each entity, without repeating every child list.
    for edge in MODEL['edges']:
        if edge['child']!=key:continue
        parent=N[edge['parent']]['label']
        maximum='one' if edge['unique'] else 'many'
        optional='0..1' if edge['optional'] else '1'
        symbol=' -- ' if edge['unique'] else ' --< '
        out.append(f"{parent}{symbol}{t['label']} ({edge['kind']}: {', '.join(edge['cols'])}; parent {optional}, child 0..{'1' if edge['unique'] else '*'})")
    for bridge in MODEL['bridges']:
        if key not in (bridge['parent'],bridge['child']):continue
        other=bridge['child'] if key==bridge['parent'] else bridge['parent']
        out.append(f"{t['label']} >--< {N[other]['label']} via {bridge['via']}")
    # Root entities still expose their main aggregate relationship.
    if not out:
        for edge in MODEL['edges']:
            if edge['parent']==key and edge['kind']=='FK':
                out.append(f"{t['label']} --{' ' if edge['unique'] else '< '}{N[edge['child']]['label']}")
    if 'source_reference_id' in t['cols']:
        out.append('Source selected by source type; no fixed FK.')
    if key=='assessment.grading_point_costs':out=['Standalone cost configuration; grading jobs retain the cost snapshot, not a cost-row FK.']
    if key=='assessment.skill_scores':out.append('feedback_revision_id has no declared FK target.')
    if key=='learning.kp_evidence':out.append('Source reference is polymorphic; mastery is derived from evidence, not a separate stored entity.')
    return '\n'.join(out) or 'No additional fixed relationship declared in the SQL scripts.'

para('This section summarizes the core business entities and their key relationships from the existing Flyway SQL scripts. Each entity is owned by the service shown below. Names are conceptual; the source map beside the diagram links them to the existing tables. This is a data model, not an assertion that every provider integration or stored capability is implemented.')
para('Notation: A --< B means one-to-many; A -- B means one-to-one (optional participation is stated explicitly); A >--< B means many-to-many through the named association. In the diagrams, a circle means zero, a bar means one and a crow\'s foot means many. PK marks an identifier; FK marks a declared same-database foreign key; Ref marks a stored identifier without a declared foreign key. Solid lines show declared FK/association links; dashed lines show logical references. A shaded repeated entity identifies a reference endpoint, not shared database ownership.')
para('Figures 4.1a–4.1n are views of one conceptual ERD. Selected key links appear in each view; the entity tables also list references that cross views or services. No cross-service link is a database foreign key. Many-to-many association records are represented by their links rather than separate core entities. Full attributes belong in IELTSPath_RTW.xlsx, Sheet 5; physical schema, indexes and migration details belong in the TDS.')
para('Source: services/<service>-service/src/main/resources/db/migration/*.sql, read in version order through the current scripts. The Content V7 removal and Library V1/V2 ownership are reflected, as are Content V16 lesson links and Learning V4/V5 practice and review records. Diagram source: part4-diagrams/IELTSPath_ERD.drawio; migration-to-entity mapping: part4-diagrams/README.md.')

def add_table(keys):
    el=deepcopy(table_template)
    rows=el.findall(qn('w:tr'))
    sample=deepcopy(rows[1])
    for row in rows[1:]:el.remove(row)
    widths=[1.15,1.7,2.75,1.0]
    grid=el.find(qn('w:tblGrid'))
    for c,w in zip(grid,widths):c.set(qn('w:w'),str(int(w*1440)))
    header=el.find(qn('w:tr'))
    trpr=header.find(qn('w:trPr'))
    if trpr is None:trpr=OxmlElement('w:trPr');header.insert(0,trpr)
    if trpr.find(qn('w:tblHeader')) is None:trpr.append(OxmlElement('w:tblHeader'))
    for ci,c in enumerate(header.findall(qn('w:tc'))):
        c.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'),str(int(widths[ci]*1440)))
    for key in keys:
        n=N[key]; row=deepcopy(sample)
        rp=row.find(qn('w:trPr'))
        if rp is None:rp=OxmlElement('w:trPr');row.insert(0,rp)
        rp.append(OxmlElement('w:cantSplit'))
        for i,(cell,value) in enumerate(zip(row.findall(qn('w:tc')),[n['label'],n['purpose'],relationship_text(key),'RTW.xlsx\nSheet 5'])):
            cell.find(qn('w:tcPr')).find(qn('w:tcW')).set(qn('w:w'),str(int(widths[i]*1440)))
            pbase=deepcopy(cell.find(qn('w:p')))
            for p in list(cell):
                if p.tag!=qn('w:tcPr'):cell.remove(p)
            for line in value.split('\n'):
                p=deepcopy(pbase)
                for r in list(p):
                    if r.tag!=qn('w:pPr'):p.remove(r)
                pp=p.find(qn('w:pPr'))
                if pp is None:pp=OxmlElement('w:pPr');p.insert(0,pp)
                for prop in ('keepNext','keepLines'):
                    old=pp.find(qn('w:'+prop))
                    if old is not None:pp.remove(old)
                r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
                size=OxmlElement('w:sz');size.set(qn('w:val'),'19');pr.append(size)
                color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color)
                if i==0:pr.append(OxmlElement('w:b'))
                r.append(pr); t=OxmlElement('w:t');t.text=line;r.append(t);p.append(r);cell.append(p)
        el.append(row)
    anchor.addprevious(el);new_elements.append(el)

for view in VIEWS:
    letter=chr(96+view['index'])
    para(f"{view['title']} ({view['service']}-service)", 'Heading 3',break_before=True)
    image=Image.open(DIR/view['png']); iw,ih=image.size
    # Keep diagrams readable and inside A4; tall views reserve a page for the figure.
    width=min(6.6,8.0*iw/ih)
    p=para();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(DIR/view['png']),width=Inches(width))
    cap=para(f"Figure 4.1{letter} — Conceptual ERD: {view['title']}. Source: IELTSPath_ERD.drawio, page {view['index']}.")
    cap.paragraph_format.space_after=Pt(8)
    for r in cap.runs:r.italic=True;r.font.size=Pt(9)
    if ih/iw>1:para('Entity summary',bold=True,break_before=True)
    add_table(view['primary'])
    para()

para('Supporting structures are not separate core entities in these views: user-role and deck membership links, content question/KP/asset links, assessment KP attribution and judgments, authentication/OAuth records, outbox records, Learning catalog projections, processed assessment versions and daily LLM usage counters. They remain part of the source scripts. Notification has no schema here. No retention, deletion policy or new data entity is introduced by this section.')

# Serialize authoring work, then copy only changed package parts back into original archive.
buf=BytesIO();doc.save(buf)
with ZipFile(BACKUP) as old,ZipFile(buf) as changed:
    old_parts={n:old.read(n) for n in old.namelist()}
    changed_parts={n:changed.read(n) for n in changed.namelist()}
    allow={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
    for n in changed_parts:
        if n.startswith('word/media/') and n not in old_parts:allow.add(n)
    with ZipFile(SRS,'w',ZIP_DEFLATED) as result:
        for info in old.infolist():result.writestr(info,changed_parts[info.filename] if info.filename in allow else old_parts[info.filename])
        for n in allow-old_parts.keys():result.writestr(n,changed_parts[n])

# Prove every original body element outside the replaced section is unchanged.
after=Document(SRS);newchildren=list(after._element.body)
newstart=next(i for i,e in enumerate(newchildren) if text(e).startswith('4.1  Core Entities'))
newend=next(i for i,e in enumerate(newchildren) if text(e).startswith('4.2  Data Constraints'))
assert [E.tostring(e) for e in children[:start+1]]==[E.tostring(e) for e in newchildren[:newstart+1]]
assert [E.tostring(e) for e in children[end:]]==[E.tostring(e) for e in newchildren[newend:]]
with ZipFile(SRS) as z:
    for name,data in old_parts.items():
        if name not in allow:assert z.read(name)==data,name

# A section-only copy is for visual QA, not a separate user deliverable.
qa=Document(SRS)
for el in list(qa._element.body):
    if el.tag!=qn('w:sectPr'):qa._element.body.remove(el)
for el in newchildren[newstart:newend]:qa._element.body.insert(len(qa._element.body)-1,deepcopy(el))
qa.save(DIR/'qa'/'section-4-1.docx')

lines=['# SRS 4.1 source map','','The ERD is derived from existing SQL migrations, read only. No SQL was executed.','',
'`IELTSPath_ERD.drawio` contains the 14 editable diagram views embedded in the SRS. PNG exports use the same SQL-derived model.','',
'Solid links: declared foreign keys or named association tables. Dashed links: stored identifier references without SQL foreign keys. Cardinalities reflect nullable and unique/primary-key declarations, not an invented minimum child count.','',
'## Core entity mapping','','| Entity | Owner | Source table | Migration |','|---|---|---|---|']
for n in N.values():lines.append(f"| {n['label']} | {n['service']} | `{n['name']}` | `{n['source']}` |")
lines += ['','## Association tables','','| Association | Endpoints | Keys |','|---|---|---|']
for b in MODEL['bridges']:lines.append(f"| `{b['via']}` | {N[b['parent']]['label']} / {N[b['child']]['label']} | `{', '.join(b['cols'])}` |")
lines += ['','`content_asset_links` belongs to either a section or a question version, not both (SQL CHECK).',
'`attempt_item_knowledge_points` and `item_result_knowledge_judgments` preserve assessment-level KP attribution/judgment; neither introduces a new canonical Knowledge Point.',
'Reference fields that are polymorphic or lack a declared target (for example source_reference_id and feedback_revision_id) are not converted into guessed foreign keys.',
'','## Reproduce','','Use the installed Codex Python runtime for `build-srs-entities.py` and `update-srs-section.py`. Set `CODEX_ARTIFACT_MODULES` to the bundled Node module directory and run `render-erd.cjs` with the bundled Node runtime. The updater preserves its original input under `qa/` for scope verification.','',
'## Migration evidence','','All migrations below were inspected for CREATE/ALTER/DROP statements. Seed records do not define new entities.','']
for s in MODEL['sources']:lines.append(f"- `{s['path']}` — SHA-256 `{s['sha256']}`")
(DIR/'README.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
(DIR/'qa'/'verification.json').write_text(json.dumps({'entities':len(N),'figures':len(VIEWS),'unchanged_outside_section':True,'preserved_package_parts':len(old_parts)-len(allow & old_parts.keys()),'updated_part':'4.1','original_sha256':hashlib.sha256(BACKUP.read_bytes()).hexdigest()},indent=2))
print(f'Updated only 4.1: {len(N)} entity rows, {len(VIEWS)} ERD views; outside-section XML preserved.')
